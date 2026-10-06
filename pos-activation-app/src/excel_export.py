"""
Generates the downloadable 8-sheet .xlsx workbook using openpyxl.

Sheet order (consolidation must be first):
    1. Division-wise Consolidation
    2-8. One sheet per division, in the fixed order from config.DIVISIONS

The consolidation sheet follows the report owner's reference proforma:
a single two-line title block (report name + period, then the activation
criteria in brackets) directly above the header row, divisions ranked by
Activation % descending, and each division row banded sea-green (meets the
config.DIVISION_ACTIVATION_CRITERIA_PCT threshold) or light pink (below it)
— the Total row keeps its own distinct styling regardless of that threshold.

Division sheets keep a simple Title -> blank -> header -> data layout (the
report owner asked for the old subtitle/reporting-date lines to be dropped),
with the existing per-office GREEN/ORANGE/RED banding unchanged.
"""

from __future__ import annotations

import io

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import config
from .report import CONSOLIDATION_HEADERS, DIVISION_SHEET_COLUMNS, DIVISION_SHEET_HEADERS

THIN = Side(style="thin", color="B7B7B7")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

HEADER_FONT = Font(name=config.FONT_NAME, size=11, bold=True, color=config.FONT_HEADER)
HEADER_FILL = PatternFill("solid", fgColor=config.FILL_HEADER)
TITLE_FONT = Font(name=config.FONT_NAME, size=13, bold=True, color="1F4E78")
TOTAL_FONT = Font(name=config.FONT_NAME, size=11, bold=True)
TOTAL_FILL = PatternFill("solid", fgColor=config.FILL_TOTAL_ROW)
BODY_FONT = Font(name=config.FONT_NAME, size=10)

STATUS_STYLE = {
    config.STATUS_GREEN: (PatternFill("solid", fgColor=config.FILL_GREEN), Font(name=config.FONT_NAME, size=10, bold=True, color=config.FONT_GREEN)),
    config.STATUS_ORANGE: (PatternFill("solid", fgColor=config.FILL_ORANGE), Font(name=config.FONT_NAME, size=10, bold=True, color=config.FONT_ORANGE)),
    config.STATUS_RED: (PatternFill("solid", fgColor=config.FILL_RED), Font(name=config.FONT_NAME, size=10, bold=True, color=config.FONT_RED)),
}

DIVISION_ROW_FILLS = {
    True: PatternFill("solid", fgColor=config.FILL_DIVISION_MEETS_CRITERIA),   # meets criteria
    False: PatternFill("solid", fgColor=config.FILL_DIVISION_BELOW_CRITERIA),  # below criteria
}

INT_FORMAT = "#,##0"
AMOUNT_FORMAT = '"\u20b9"#,##0.00'
PERCENT_FORMAT = "0.00%"


def consolidation_title_text(period_start_str: str, period_end_str: str) -> str:
    """Builds the exact two-line quoted title from the report owner's
    reference proforma, with dates and thresholds pulled from config/input
    rather than hardcoded, e.g.:

        "SBI POS Machine Activation Status (Period: 01.09.2026 to 30.09.2026)
        [Monthly Criteria for activation: No. of Transactions >= 10 /
        Transaction amt >= Rs.5000/- Division activation criteria: 50%]"
    """
    pct = int(round(config.DIVISION_ACTIVATION_CRITERIA_PCT * 100))
    line1 = f'"SBI POS Machine Activation Status (Period: {period_start_str} to {period_end_str})'
    line2 = (
        f"[Monthly Criteria for activation: No. of Transactions >= {config.COUNT_THRESHOLD_PER_DEVICE} / "
        f'Transaction amt >= Rs.{config.AMOUNT_THRESHOLD_PER_DEVICE}/- '
        f'Division activation criteria: {pct}%]"'
    )
    return line1 + "\n" + line2


def _write_consolidation_title_block(ws: Worksheet, title_text: str, num_cols: int) -> int:
    """Single merged, wrapped, 2-line title directly above the header row —
    no separate subtitle/date rows. Returns the header row index."""
    ws["A1"] = title_text
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(num_cols, 4))
    ws.row_dimensions[1].height = 42
    return 2  # header row


def _write_division_title_block(ws: Worksheet, title: str, num_cols: int) -> int:
    """Title -> blank -> header (the old subtitle/reporting-date rows have
    been removed per the report owner's instruction)."""
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(num_cols, 4))
    return 3  # header row (row 2 left blank as a visual gap)


def _autosize_columns(ws: Worksheet, headers: list, sample_rows: pd.DataFrame, keys: list,
                       min_width=8, max_width=42):
    for idx, (key, header) in enumerate(zip(keys, headers), start=1):
        max_len = len(str(header))
        if len(sample_rows) > 0 and key in sample_rows.columns:
            col_len = sample_rows[key].astype(str).map(len).max()
            max_len = max(max_len, int(col_len))
        width = min(max(max_len + 3, min_width), max_width)
        ws.column_dimensions[get_column_letter(idx)].width = width


def _write_header_row(ws: Worksheet, row_idx: int, headers: list):
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = CELL_BORDER
    ws.row_dimensions[row_idx].height = 30


def write_consolidation_sheet(wb: Workbook, consolidation_df: pd.DataFrame,
                               period_start_str: str, period_end_str: str) -> Worksheet:
    ws = wb.create_sheet(config.CONSOLIDATION_SHEET_NAME)
    keys = list(CONSOLIDATION_HEADERS.keys())
    headers = list(CONSOLIDATION_HEADERS.values())

    title_text = consolidation_title_text(period_start_str, period_end_str)
    header_row = _write_consolidation_title_block(ws, title_text, len(headers))
    _write_header_row(ws, header_row, headers)

    first_data_row = header_row + 1
    for i, (_, row) in enumerate(consolidation_df.iterrows()):
        r = first_data_row + i
        is_total = str(row["division"]).strip().lower() == "total"
        meets_criteria = row["activation_pct"] >= config.DIVISION_ACTIVATION_CRITERIA_PCT
        for c, key in enumerate(keys, start=1):
            value = row[key]
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = CELL_BORDER
            if is_total:
                cell.font = TOTAL_FONT
                cell.fill = TOTAL_FILL
            else:
                cell.font = BODY_FONT
                cell.fill = DIVISION_ROW_FILLS[meets_criteria]
            if key == "activation_pct":
                cell.number_format = PERCENT_FORMAT
                cell.alignment = Alignment(horizontal="center")
            elif key in ("devices_supplied", "devices_activated", "sl_no"):
                cell.number_format = INT_FORMAT
                cell.alignment = Alignment(horizontal="center")
            elif key == "division":
                cell.alignment = Alignment(horizontal="left")
            else:
                cell.alignment = Alignment(horizontal="center")

    last_row = first_data_row + len(consolidation_df) - 1
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    ws.freeze_panes = f"A{first_data_row}"

    _autosize_columns(ws, headers, consolidation_df, keys)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def write_division_sheet(wb: Workbook, division_name: str, df: pd.DataFrame) -> Worksheet:
    ws = wb.create_sheet(division_name)
    keys = DIVISION_SHEET_COLUMNS
    headers = [DIVISION_SHEET_HEADERS[k] for k in keys]

    header_row = _write_division_title_block(ws, f"POS Activation Status \u2014 {division_name} Division", len(headers))
    _write_header_row(ws, header_row, headers)

    first_data_row = header_row + 1
    for i, (_, row) in enumerate(df.iterrows()):
        r = first_data_row + i
        fill, status_font = STATUS_STYLE[row["activation_status"]]
        for c, key in enumerate(keys, start=1):
            value = row[key]
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = CELL_BORDER
            cell.fill = fill
            if key == "activation_status":
                cell.font = status_font
                cell.alignment = Alignment(horizontal="center")
            else:
                cell.font = BODY_FONT
                cell.alignment = Alignment(horizontal="center") if key != "office_name" else Alignment(horizontal="left")

            if key in ("sbipos_card_amount", "sbipos_bharatqr_amount", "combined_amount", "amount_threshold"):
                cell.number_format = AMOUNT_FORMAT
            elif key in ("sl_no", "devices_supplied", "sbipos_card_count", "sbipos_bharatqr_count",
                         "combined_count", "count_threshold", "sort_group"):
                cell.number_format = INT_FORMAT
            elif key == "office_id":
                cell.number_format = "@"  # keep as text, preserve any leading zeros

    last_row = first_data_row + len(df) - 1 if len(df) else first_data_row
    if len(df):
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{last_row}"
    ws.freeze_panes = f"A{first_data_row}"

    _autosize_columns(ws, headers, df, keys)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def build_workbook(consolidation_df: pd.DataFrame, division_sheets: dict,
                    period_start_str: str, period_end_str: str) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank sheet

    write_consolidation_sheet(wb, consolidation_df, period_start_str, period_end_str)
    for division in config.DIVISIONS:
        write_division_sheet(wb, division, division_sheets[division])

    assert wb.sheetnames == config.WORKBOOK_SHEET_ORDER, (
        f"Sheet order mismatch: {wb.sheetnames} != {config.WORKBOOK_SHEET_ORDER}"
    )
    assert len(wb.sheetnames) == 8, f"Expected exactly 8 sheets, got {len(wb.sheetnames)}"
    return wb


def workbook_to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
